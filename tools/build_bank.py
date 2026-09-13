#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
题库导入工具：把 sources/ 下抓取的第三方开源题库转换为网站 data/*.json 格式。

用法：python tools/build_bank.py
- 读取 sources/ 目录中的原始文件（不入 git，见 .gitignore）
- 统一清洗：去 HTML 标签、去题号、推断/归一化题型、选项打乱（答案跟随）
- 同行业按题干去重（已入库的旧题优先保留）
- 输出合并到 data/*.json（按分类 id 幂等合并），打印统计报告
"""
import json
import os
import re
import sys
import glob
import random
import collections

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "sources")
LETTERS = "ABCDEFG"

TAG_RE = re.compile(r"<[^>]+>")
NUM_RE = re.compile(r"^\s*\(?\d{1,3}[.、．）)]\s*")


def strip_html(s):
    if not s:
        return ""
    s = TAG_RE.sub("", str(s))
    return normalize(s)


def normalize(s):
    s = str(s).replace("\xa0", " ").replace("\u3000", " ").replace("\ufeff", "")
    s = re.sub(r"<br\s*/?>", " ", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def clean_stem(s):
    s = strip_html(s)
    return NUM_RE.sub("", s).strip()


# ---------------- 各数据源解析器：yield dict(type, q, options, answer, explain, subject) ----------------

def parse_diangongtiku():
    """PLSshenyun/diangongtiku 电工题库 CSV（GBK）：id,题干,4选项,答案字母 / id,题干,T|F"""
    import csv
    p = os.path.join(SRC, "diangongtiku_questions_choose.csv")
    if os.path.exists(p):
        for r in csv.reader(open(p, encoding="gbk", errors="replace")):
            if len(r) >= 7 and r[-1].strip() in "ABCD" and len(r[1]) > 8:
                yield dict(type="single", q=clean_stem(r[1]), options=[strip_html(x) for x in r[2:6]],
                           answer="ABCD".index(r[-1].strip()), explain="", subject=None)
    p = os.path.join(SRC, "diangongtiku_questions_judge.csv")
    if os.path.exists(p):
        for r in csv.reader(open(p, encoding="gbk", errors="replace")):
            if len(r) >= 3 and r[-1].strip() in "TF" and len(r[1]) > 8:
                yield dict(type="judge", q=clean_stem(r[1]), options=[], answer=r[-1].strip() == "T",
                           explain="", subject=None)


def parse_diyadiangong():
    """PTA00/diyadiangong 低压电工 判断.txt：id#题干$0|1（0=正确,1=错误）"""
    p = os.path.join(SRC, "diyadiangong_低压电工练习_数据采集原文件_判断.txt")
    if os.path.exists(p):
        for line in open(p, encoding="utf-8", errors="replace"):
            line = line.strip()
            if "#" not in line or "$" not in line:
                continue
            stem, _, flag = line.rpartition("$")
            stem = stem.split("#", 1)[1] if "#" in stem else stem
            if len(stem) > 8 and flag in ("0", "1"):
                yield dict(type="judge", q=clean_stem(stem), options=[], answer=flag == "0", explain="", subject=None)


def parse_erjian():
    """xiaoxiunique/ac 二级建造师 questions.json：{list:[{type:'1'|'2', stem(html), options(html), answer:[i], analysis}]}"""
    p = os.path.join(SRC, "ac_src_pages_questions.json")
    if not os.path.exists(p):
        return
    d = json.load(open(p, encoding="utf-8"))["list"]
    for q in d:
        if q.get("type") not in ("1", "2"):
            continue  # type 5 等为图片/主观题
        stem = strip_html(q.get("stem", ""))
        if "<img" in q.get("stem", "") or len(stem) < 10:
            continue
        opts = [strip_html(o) for o in q.get("options", [])]
        if len(opts) < 2 or any(len(o) < 1 for o in opts):
            continue
        ans = [a for a in q.get("answer", []) if isinstance(a, int) and 0 <= a < len(opts)]
        if not ans:
            continue
        yield dict(type="single" if q["type"] == "1" else "multi", q=stem, options=opts,
                   answer=ans[0] if q["type"] == "1" else sorted(set(ans)),
                   explain=strip_html(q.get("analysis", "")), subject=None)


def parse_fakao():
    """chenjackie178/legal-exam-assistant 法考真题 questions_all.json"""
    p = os.path.join(SRC, "falv_fakao_questions_all.json")
    if not os.path.exists(p):
        return
    d = json.load(open(p, encoding="utf-8"))["questions"]
    for q in d:
        qt = q.get("question_type", "")
        if qt not in ("single_choice", "multiple_choice"):
            continue
        opts = list(q.get("options", {}).values())
        opts = [strip_html(o) for o in opts]
        if len(opts) < 2:
            continue
        ans = str(q.get("answer", "")).strip().upper()
        idx = [LETTERS.index(c) for c in ans if c in LETTERS[:len(opts)]]
        if not idx or len(set(idx)) != len(idx):
            continue
        yield dict(type="single" if qt == "single_choice" else "multi", q=clean_stem(q.get("content", "")),
                   options=opts, answer=idx[0] if qt == "single_choice" else sorted(set(idx)),
                   explain=strip_html(q.get("analysis", "")), subject=None)


def parse_fadao():
    """andesiwangzhiyi-alt/fadao 法考分科题库 js/questions.js|questions2.js（MIT）
    文件形如 const QUESTION_BANK = { "法理学": [ {stem, options, answer:[i], analysis}, ... ] }
    键名不带引号且字符串内含 "A:" 之类文本，用 json5 宽松解析。"""
    import json5
    for fname in ("falv_fadao_questions1.js", "falv_fadao_questions2.js"):
        p = os.path.join(SRC, fname)
        if not os.path.exists(p):
            continue
        js = open(p, encoding="utf-8", errors="replace").read()
        m = re.search(r"=\s*\{", js)
        if not m:
            continue
        body = js[m.start() + 1:]
        body = body[: body.rfind("}") + 1]
        try:
            bank = json5.loads(body)
        except Exception as e:
            print("  [fadao] %s 解析失败: %s" % (fname, str(e)[:80]))
            continue
        for subj, items in bank.items():
            if not isinstance(items, list):
                continue
            for q in items:
                if not isinstance(q, dict):
                    continue
                opts = [strip_html(o) for o in q.get("options", []) if strip_html(o)]
                ans = [a for a in (q.get("answer") or []) if isinstance(a, int) and 0 <= a < len(opts)]
                stem = clean_stem(q.get("stem", ""))
                if len(opts) < 2 or len(stem) < 10 or not ans:
                    continue
                yield dict(type="multi" if len(ans) > 1 else "single", q=stem, options=opts,
                           answer=ans[0] if len(ans) == 1 else sorted(set(ans)),
                           explain=strip_html(q.get("analysis", "")), subject=normalize(subj))


def parse_zhongyi():
    """xyfan57/zhongyi-zhiye-shuati 中医执业医师 questions.json：{stem, options:[[字母,文本]], answer:'A', explanation, subject}"""
    p = os.path.join(SRC, "yixue_zhongyi_questions.json")
    if not os.path.exists(p):
        return
    for q in json.load(open(p, encoding="utf-8")):
        pairs = q.get("options", [])
        opts = [strip_html(t) for _, t in pairs]
        ans = str(q.get("answer", "")).strip().upper()
        if len(opts) < 2 or ans not in LETTERS[:len(opts)] or len(clean_stem(q.get("stem", ""))) < 8:
            continue
        yield dict(type="single", q=clean_stem(q.get("stem", "")), options=opts,
                   answer=LETTERS.index(ans), explain=strip_html(q.get("explanation", "")),
                   subject=normalize(q.get("subject", "")) or None)


def parse_zhongji_kuaiji():
    """Lmrean/IAT-exam 中级会计 bank.json：{subject 财管|实务, type single|multi|judge|calc|comprehensive, stem, options{A:D}, answer, analysis}"""
    p = os.path.join(SRC, "caikuai_zhongji_bank.json")
    if not os.path.exists(p):
        return
    for q in json.load(open(p, encoding="utf-8")):
        t = q.get("type")
        if t not in ("single", "multi", "judge"):
            continue  # calc/comprehensive 为主观题
        stem = clean_stem(q.get("stem", ""))
        if len(stem) < 10:
            continue
        opts = [strip_html(v) for k, v in sorted((q.get("options") or {}).items())]
        if t in ("single", "multi"):
            if len(opts) < 2 or any(not o for o in opts):
                continue
            ans = str(q.get("answer", "")).strip().upper()
            idx = [LETTERS.index(c) for c in ans if c in LETTERS[:len(opts)]]
            if not idx or len(set(idx)) != len(idx):
                continue
            yield dict(type=t, q=stem, options=opts, answer=idx[0] if t == "single" else sorted(set(idx)),
                       explain=strip_html(q.get("analysis", "")), subject=normalize(q.get("subject", "")))
        else:
            a = str(q.get("answer", "")).strip().upper()
            if a in ("T", "正确", "对", "Y", "√"):
                ans = True
            elif a in ("F", "错误", "错", "N", "×", "X"):
                ans = False
            else:
                continue
            yield dict(type="judge", q=stem, options=[], answer=ans,
                       explain=strip_html(q.get("analysis", "")), subject=normalize(q.get("subject", "")))


def parse_c3_anquan():
    """il565003788-rgb/c3-safety-practice 建筑施工安全题库：{content, option_a..e, correct_answer:'A'|'ABD', explanation, type}"""
    p = os.path.join(SRC, "anquan_c3_questions.json")
    if not os.path.exists(p):
        return
    for q in json.load(open(p, encoding="utf-8"))["questions"]:
        t = q.get("type")
        opts = [strip_html(q.get("option_" + c, "")) for c in "abcde"]
        opts = [o for o in opts if o]
        stem = clean_stem(q.get("content", ""))
        if len(opts) < 2 or len(stem) < 8:
            continue
        ans = str(q.get("correct_answer", "")).strip().upper()
        idx = [LETTERS.index(c) for c in ans if c in LETTERS[:len(opts)]]
        if not idx or len(set(idx)) != len(idx):
            continue
        yield dict(type="judge" if t == "judge" else ("multi" if t == "multiple" else "single"),
                   q=stem, options=opts if t != "judge" else [],
                   answer=(idx[0] if t != "multiple" and len(idx) == 1 else sorted(set(idx))) if t != "judge"
                   else (idx and ans in ("A", "T", "对", "正确")),
                   explain=strip_html(q.get("explanation", "")), subject=None)


def parse_xxqg():
    """imutum/XXQG_TiKu 学习强国挑战答题：{“题干|选项1|选项2|...”: 正确选项文本}"""
    p = os.path.join(SRC, "changshi_xxqg_tiku.json")
    if not os.path.exists(p):
        return
    d = json.load(open(p, encoding="utf-8"))
    for k, v in d.items():
        parts = [normalize(x) for x in k.split("|")]
        stem = NUM_RE.sub("", parts[0]).strip()
        stem = re.sub(r"来源[:：].*$", "", stem).rstrip("。，,；;　 ").strip()
        opts = [x for x in parts[1:] if x]
        val = normalize(v)
        if len(stem) < 8 or len(opts) < 2:
            continue
        if set(o for o in opts) >= {"正确", "错误"} and val in ("正确", "错误"):
            yield dict(type="judge", q=stem, options=[], answer=val == "正确", explain="", subject=None)
            continue
        # 找正确选项文本对应的下标（模糊匹配：完全相等或互相包含）
        idx = None
        for i, o in enumerate(opts):
            if o == val or (len(o) > 3 and len(val) > 3 and (o in val or val in o)):
                idx = i
                break
        if idx is None:
            continue
        yield dict(type="single", q=stem, options=opts, answer=idx, explain="", subject=None)


def extract_after(text, pos):
    """从 pos 处跳过空白，提取从 { 或 [ 开始的括号配对完整片段（字符串感知），返回 (字符串, 开括号) 或 (None, None)。"""
    i, n = pos, len(text)
    while i < n and text[i] in " \n\r\t":
        i += 1
    if i >= n or text[i] not in "{[":
        return None, None
    open_ch = text[i]
    close_ch = "}" if open_ch == "{" else "]"
    depth, j, in_str, esc = 0, i, False, False
    while j < n:
        c = text[j]
        if in_str:
            if esc:
                esc = False
            elif c == "\\":
                esc = True
            elif c == in_str:
                in_str = False
        else:
            if c in "\"'`":
                in_str = c
            elif c == open_ch:
                depth += 1
            elif c == close_ch:
                depth -= 1
                if depth == 0:
                    return text[i: j + 1], open_ch
        j += 1
    return None, None


def parse_fire():
    """ybd0612/fire-operator-study 消防设施操作员中级模拟卷 t1-t4.txt
    纯文本格式：题号行、题干行、A/B/C/D 单独行+内容行、“正确答案：X”、“参考解析：...”"""
    pats_sec = re.compile(r"^(单选题|多选题|判断题)（共")
    pat_num = re.compile(r"^(\d{1,3})\.$")
    pat_opt = re.compile(r"^([A-F])$")
    pat_ans = re.compile(r"正确答案[:：]\s*([A-F对错正确错误]+)")
    for i in "1234":
        p = os.path.join(SRC, "xiaofang_t%s.txt" % i)
        if not os.path.exists(p):
            continue
        lines = [l.rstrip() for l in open(p, encoding="utf-8", errors="replace")]
        sec, cur = None, None

        def flush(q):
            if not q or not q.get("stem"):
                return
            t, stem, opts, ans, exp = q["sec"], q["stem"], q["opts"], q["ans"], q["exp"]
            stem = clean_stem(stem)
            if len(stem) < 8:
                return
            if t == "judge":
                if ans in ("正确", "对", "A"):
                    yield_q = dict(type="judge", q=stem, options=[], answer=True, explain=exp, subject=None)
                elif ans in ("错误", "错", "B"):
                    yield_q = dict(type="judge", q=stem, options=[], answer=False, explain=exp, subject=None)
                else:
                    return
                yield_q and out.append(yield_q)
            else:
                if len(opts) < 2:
                    return
                idx = [LETTERS.index(c) for c in ans if c in LETTERS[:len(opts)]]
                if not idx:
                    return
                out.append(dict(type="multi" if len(idx) > 1 else "single", q=stem, options=opts[:6],
                                answer=idx[0] if len(idx) == 1 else sorted(set(idx)), explain=exp, subject=None))

        out = []
        for ln in lines:
            s = ln.strip()
            if pats_sec.match(s):
                sec = s[:3]
                cur = None
                continue
            if pat_num.match(s):
                if cur:
                    flush(cur)
                cur = dict(sec=sec, stem="", opts=[], cur_opt=None, ans="", exp="", in_exp=False)
                continue
            if cur is None:
                continue
            m = pat_ans.search(s)
            if m:
                cur["ans"] = m.group(1)
                cur["in_exp"] = False
                continue
            if s.startswith("参考解析"):
                cur["in_exp"] = True
                tail = s.split("：", 1)[-1].split(":", 1)[-1].strip()
                if tail:
                    cur["exp"] += tail + " "
                continue
            if s in ("回答正确", "回答错误") or s.startswith("我的答案"):
                continue
            if sec == "judge" and s in ("正确", "错误", "对", "错"):
                continue
            mo = pat_opt.match(s)
            if mo and sec != "judge":
                cur["cur_opt"] = mo.group(1)
                cur["opts"].append("")
                continue
            if not s:
                continue
            if cur["in_exp"]:
                cur["exp"] += s + " "
            elif cur["cur_opt"] is not None and sec != "judge":
                if cur["opts"]:
                    cur["opts"][-1] = (cur["opts"][-1] + " " + s).strip()
            else:
                cur["stem"] += s + " "
        if cur:
            flush(cur)
        for r in out:
            yield r


def parse_grammar():
    """XiaoRui114514/english-grammar-lab 英语语法题库 data/bank-*.js
    window.__GRAMMAR_LAB__.<id> = { title, questions: [{type:'choice'|'input', ...}] }"""
    import json5
    for p in sorted(glob.glob(os.path.join(SRC, "english_data_bank-*.js"))):
        js = open(p, encoding="utf-8", errors="replace").read()
        m = re.search(r"__GRAMMAR_LAB__\.\w+\s*=\s*", js)
        if not m:
            continue
        obj, _ = extract_after(js, m.end())
        if not obj:
            continue
        try:
            d = json5.loads(obj)
        except Exception as e:
            print("  [grammar] %s 解析失败: %s" % (os.path.basename(p), str(e)[:60]))
            continue
        title = normalize(d.get("title", ""))

        def exp_text(e):
            if isinstance(e, dict):
                return normalize(" ".join(str(v) for v in e.values() if v))
            return normalize(str(e or ""))

        for q in d.get("questions", []):
            qt = q.get("type")
            if qt == "choice":
                opts = [normalize(str(o)) for o in (q.get("options") or [])]
                opts = [o for o in opts if o]
                ans = q.get("answerIndex")
                stem = normalize(str(q.get("question") or ""))
                if len(opts) < 2 or len(stem) < 4 or not isinstance(ans, int) or not (0 <= ans < len(opts)):
                    continue
                yield dict(type="single", q=stem, options=opts, answer=ans,
                           explain=exp_text(q.get("explanation")), subject=title or None)
            elif qt == "input":
                stem = normalize(str(q.get("question") or ""))
                answers = [normalize(str(a)) for a in (q.get("accepted") or []) if a]
                if not answers and q.get("answerText"):
                    answers = [normalize(str(q["answerText"]))]
                answers = [a for a in answers if a]
                if len(stem) < 6 or not answers:
                    continue
                yield dict(type="blank", q=stem, answers=answers[:4],
                           explain=exp_text(q.get("explanation")), subject=title or None)


def parse_jiaozi():
    """ICEfrost777/jiaozi-fuxi-zhushou 教资单文件应用：QUESTIONS_S1/S2/S3 内嵌题库"""
    import json5
    p = os.path.join(SRC, "jiaoshi_jiaozi_index.html")
    if not os.path.exists(p):
        return
    h = open(p, encoding="utf-8", errors="replace").read()
    names = {"QUESTIONS_S1": "科目一·综合素质", "QUESTIONS_S2": "科目二·教育知识与能力", "QUESTIONS_S3": "科目三·学科知识（物理）"}
    for var, cat in names.items():
        m = re.search(var + r"\s*=\s*", h)
        if not m:
            continue
        obj, kind = extract_after(h, m.end())
        if not obj:
            continue
        try:
            items = json5.loads(obj)
        except Exception as e:
            print("  [jiaozi] %s 解析失败: %s" % (var, str(e)[:60]))
            continue
        if not isinstance(items, list):
            continue
        for q in items:
            if not isinstance(q, dict):
                continue
            stem = clean_stem(str(q.get("stem") or q.get("q") or ""))
            opts = [normalize(str(o)) for o in (q.get("options") or [])]
            opts = [o for o in opts if o]
            ans = q.get("ans")
            exp = normalize(str(q.get("exp") or ""))
            if len(opts) < 2 or len(stem) < 8:
                continue
            if q.get("type") == "judge":
                if isinstance(ans, int) and 0 <= ans < len(opts):
                    yield dict(type="judge", q=stem, options=[], answer=(opts[ans] == "正确" if ans < len(opts) else False),
                               explain=exp, subject=cat)
                continue
            if isinstance(ans, list):
                aidx = [a for a in ans if isinstance(a, int) and 0 <= a < len(opts)]
                if not aidx:
                    continue
                yield dict(type="multi" if len(aidx) > 1 else "single", q=stem, options=opts,
                           answer=aidx[0] if len(aidx) == 1 else sorted(set(aidx)), explain=exp, subject=cat)
            elif isinstance(ans, int) and 0 <= ans < len(opts):
                yield dict(type="single", q=stem, options=opts, answer=ans, explain=exp, subject=cat)


def parse_ntce():
    """sodalevy/ntce_exam 教师资格证（综合素质）questions.json：{type, question, options{A:..}, answer:'C', explanation}"""
    p = os.path.join(SRC, "jiaoshi_ntce_questions.json")
    if not os.path.exists(p):
        return
    for q in json.load(open(p, encoding="utf-8")).get("questions", []):
        t = q.get("type", "")
        opts = [strip_html(v) for k, v in sorted((q.get("options") or {}).items())]
        opts = [o for o in opts if o]
        stem = clean_stem(q.get("question", ""))
        if len(opts) < 2 or len(stem) < 8:
            continue
        ans = str(q.get("answer", "")).strip().upper()
        idx = [LETTERS.index(c) for c in ans if c in LETTERS[:len(opts)]]
        if not idx or len(set(idx)) != len(idx):
            continue
        if t == "multiple_choice" or len(idx) > 1:
            yield dict(type="multi", q=stem, options=opts, answer=sorted(set(idx)),
                       explain=strip_html(q.get("explanation", "")), subject=None)
        else:
            yield dict(type="single", q=stem, options=opts, answer=idx[0],
                       explain=strip_html(q.get("explanation", "")), subject=None)


def parse_futures():
    """zhanghanlun2023/futures-exam-2026 期货从业题库（先经 node 转为 jinrong_futures_all.json）"""
    p = os.path.join(SRC, "jinrong_futures_all.json")
    if not os.path.exists(p):
        return
    for q in json.load(open(p, encoding="utf-8")):
        t = q.get("type")
        stem = clean_stem(q.get("q", ""))
        opts = [normalize(str(o)) for o in (q.get("options") or [])]
        opts = [o for o in opts if o]
        if len(stem) < 8:
            continue
        raw = q.get("answer") or []
        raw = [str(a).strip().upper() for a in raw] if isinstance(raw, list) else [str(raw).strip().upper()]
        if t == "judge":
            txt = "".join(raw)
            # 优先按选项字母映射（选项通常是 [正确, 错误]）
            if opts and len(opts) >= 2 and raw and raw[0] in LETTERS[:len(opts)]:
                chosen = opts[LETTERS.index(raw[0])]
                if chosen in ("正确", "错误", "对", "错"):
                    yield dict(type="judge", q=stem, options=[], answer=chosen in ("正确", "对"),
                               explain=normalize(str(q.get("explain", ""))), subject=None)
                    continue
            if any(x in txt for x in ("正确", "对", "T")):
                yield dict(type="judge", q=stem, options=[], answer=True, explain=normalize(str(q.get("explain", ""))), subject=None)
            elif any(x in txt for x in ("错误", "错", "F")):
                yield dict(type="judge", q=stem, options=[], answer=False, explain=normalize(str(q.get("explain", ""))), subject=None)
            continue
        if len(opts) < 2:
            continue
        idx = [LETTERS.index(c) for c in raw if c in LETTERS[:len(opts)]]
        if not idx or len(set(idx)) != len(idx):
            continue
        yield dict(type="multi" if t == "multiple" or len(idx) > 1 else "single", q=stem, options=opts,
                   answer=idx[0] if len(idx) == 1 else sorted(set(idx)),
                   explain=normalize(str(q.get("explain", ""))), subject=normalize(str(q.get("chapter", ""))) or None)


def html_unescape(s):
    import html as _h
    return _h.unescape(str(s))


def parse_net_html():
    """zkjx/computer-network 计算机网络章节题库（THEOL 平台保存的测验结果页，GBK）。
    每题一个 <table class="infotable">：hidden input `<n>_content` 存题干，radio/checkbox 为选项，
    “[参考答案]”行为正确答案（选项文本）。"""
    pat_opt = re.compile(r'<input type="(radio|checkbox)"[^>]*name="answer_\d+"[^>]*>\s*([^<\r\n]+)')
    for p in sorted(glob.glob(os.path.join(SRC, "jisuanji_net_*.html"))):
        h = open(p, "rb").read().decode("gbk", errors="replace")
        for chunk in h.split('<table cellpadding="0" cellspacing="0" class="infotable">')[1:]:
            mc = re.search(r'name="\d+_content"[^>]*value="([^"]*)"', chunk)
            if not mc:
                continue
            stem = clean_stem(html_unescape(html_unescape(mc.group(1))))
            opts = [(t, re.sub(r"\s+", " ", html_unescape(html_unescape(txt)).replace("&nbsp;", " ")).strip())
                    for t, txt in pat_opt.findall(chunk)]
            opts = [(t, txt) for t, txt in opts if txt]
            mr = re.search(r"\[参考答案\]\s*([^<]+)", chunk)
            if len(opts) < 2 or len(stem) < 8 or not mr:
                continue
            ref = [x for x in re.split(r"[\s、,，]+", mr.group(1).strip()) if x]
            idx = []
            for a in ref:
                for i, (_, txt) in enumerate(opts):
                    if a == txt or (len(a) > 1 and a in txt):
                        if i not in idx:
                            idx.append(i)
                        break
            if not idx:
                continue
            multi = any(t == "checkbox" for t, _ in opts) or len(idx) > 1
            yield dict(type="multi" if multi else "single", q=stem,
                       options=[txt for _, txt in opts],
                       answer=idx[0] if len(idx) == 1 else sorted(idx), explain="", subject=None)


def parse_ncre_xlsx():
    """dengcao/ncre 一级 MS Office 选择题 xlsx（“自己打印” sheet：序号/题目/A-D/答案/解析）"""
    p = os.path.join(SRC, "jisuanji_ncre1_office.xlsx")
    if not os.path.exists(p):
        return
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True)
    if "自己打印" not in wb.sheetnames:
        return
    ws = wb["自己打印"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        vals = ["" if v is None else str(v).strip() for v in row]
        if len(vals) < 8:
            continue
        stem, opts, ans, exp = vals[1], vals[2:6], vals[6], vals[7]
        stem = clean_stem(stem)
        opts = [re.sub(r"^[A-D][).、．]?\s*", "", normalize(o)) for o in opts]
        opts = [o for o in opts if o]
        if len(opts) < 2 or len(stem) < 8:
            continue
        ans = ans.upper().replace("，", "").replace(" ", "")
        idx = [LETTERS.index(c) for c in ans if c in LETTERS[:len(opts)]]
        if not idx or len(set(idx)) != len(idx):
            continue
        yield dict(type="multi" if len(idx) > 1 else "single", q=stem, options=opts,
                   answer=idx[0] if len(idx) == 1 else sorted(set(idx)),
                   explain=normalize(exp), subject=None)


# ---------------- 行业/分类映射 ----------------
# 每个源路由到 (行业id, 分类id, 分类名)；按 subject 细分的在名称中体现
def route(rec):
    t, subj = rec["type"], rec.get("subject")
    src = rec["_src"]
    if src in ("diangongtiku", "diyadiangong"):
        return ("dianqi", "dgzk", "电工进网作业题库")
    if src == "erjian":
        return ("jianzhu", "ejzk", "二级建造师题库")
    if src == "c3":
        return ("jianzhu", "jzaqzk", "建筑施工安全题库")
    if src == "fakao":
        return ("falv", "fkzk", "法律职业资格考试真题")
    if src == "fadao":
        return ("falv", "fkfk", "法考分科练习")
    if src == "zhongyi":
        return ("yixue", "zy_" + (subj or "zonghe"), subj or "中医综合")
    if src == "kuaiji":
        return ("caikuai", "zjkj_" + (subj or "zonghe"), "中级会计·" + (subj or "综合"))
    if src == "xxqg":
        return ("gongwuyuan", "xxqg", "公共基础·挑战答题")
    if src == "fire":
        return ("xiaofang", "xfczy", "消防设施操作员（中级）")
    if src == "grammar":
        return ("fanyi", "yyjf", "英语基础·语法词汇")
    if src == "jiaozi":
        ids = {"科目一·综合素质": "jz_km1", "科目二·教育知识与能力": "jz_km2", "科目三·学科知识（物理）": "jz_km3"}
        return ("jiaoshi", ids.get(subj, "jz_zh"), subj or "教师资格笔试")
    if src == "ntce":
        return ("jiaoshi", "ntce", "教资笔试·综合素质题库")
    if src == "futures":
        return ("jinrong", "qhcy", "期货从业题库")
    if src == "net_html":
        return ("jisuanji", "netjk", "计算机网络章节题库")
    if src == "ncre_xlsx":
        return ("jisuanji", "ncre1", "计算机一级·选择题题库")
    return None


def route_kw(rec):
    """学习强国题目按关键词二次分流到相关行业的常识类分类（按顺序首个命中生效）"""
    q = rec["q"]
    if re.search(r"消防|火灾|灭火|疏散|逃生|救火|防火", q):
        return ("xiaofang", "xfcs", "消防安全常识")
    if re.search(r"危险化学品|化学品|爆炸极限|粉尘|职业中毒|毒物|泄漏|易燃易爆|化工|压力容器|防毒面具|安全生产许可", q):
        return ("huagong", "aqcs", "化工安全常识")
    if re.search(r"心理|情绪|抑郁|焦虑|催眠|人格障碍|心理咨询|心理健康|性格", q):
        return ("xinli", "xlcs", "心理健康常识")
    if re.search(r"计算机|软件|互联网|人工智能|操作系统|CPU|内存", q):
        return ("jisuanji", "jswk", "信息科技常识")
    if re.search(r"货币|银行|利率|证券|基金|保险|股票|贷款|汇率|信用卡|存款|金融", q):
        return ("jinrong", "jrcs", "金融基础常识")
    if re.search(r"教育|教师|学校|高考|课堂", q):
        return ("jiaoshi", "jycs", "教育常识")
    return None


# build_bank.py 托管的分类 id（--fresh 时先移除再重建，保证可重复全量重建）
MANAGED_EXACT = {"dgzk", "ejzk", "jzaqzk", "fkzk", "fkfk", "xxqg", "xfcs", "aqcs", "xfczy",
                 "yyjf", "jz_km1", "jz_km2", "jz_km3", "ntce", "qhcy", "netjk", "ncre1",
                 "xlcs", "jrcs", "jycs", "jswk"}
MANAGED_PREFIX = ("zy_", "zjkj_")


INDUSTRY_ORDER = None  # 读取现有文件顺序


def main():
    sources = [
        ("diangongtiku", parse_diangongtiku),
        ("diyadiangong", parse_diyadiangong),
        ("erjian", parse_erjian),
        ("fakao", parse_fakao),
        ("fadao", parse_fadao),
        ("zhongyi", parse_zhongyi),
        ("kuaiji", parse_zhongji_kuaiji),
        ("c3", parse_c3_anquan),
        ("xxqg", parse_xxqg),
        ("fire", parse_fire),
        ("grammar", parse_grammar),
        ("jiaozi", parse_jiaozi),
        ("ntce", parse_ntce),
        ("futures", parse_futures),
        ("net_html", parse_net_html),
        ("ncre_xlsx", parse_ncre_xlsx),
    ]
    # 行业元信息来自现有 data/*.json（跳过 manifest.json 等非行业文件）
    industries = {}
    for f in sorted(os.listdir(os.path.join(ROOT, "data"))):
        if f.endswith(".json"):
            d = json.load(open(os.path.join(ROOT, "data", f), encoding="utf-8"))
            if isinstance(d, dict) and "id" in d and "categories" in d:
                industries[d["id"]] = d

    # --fresh：先移除托管分类，实现全量重建（题目从 sources/ 重新解析）
    if "--fresh" in sys.argv:
        n_removed = 0
        for d in industries.values():
            before = len(d.get("categories", []))
            d["categories"] = [c for c in d.get("categories", [])
                               if not (c["id"] in MANAGED_EXACT or str(c["id"]).startswith(MANAGED_PREFIX))]
            n_removed += before - len(d["categories"])
        print("--fresh: 移除托管分类 %d 个，开始全量重建" % n_removed)

    # 收集
    bucket = collections.defaultdict(list)  # (indId, catId) -> [rec]
    stats = collections.Counter()
    for name, fn in sources:
        n_in, n_ok = 0, 0
        for rec in fn():
            n_in += 1
            rec["_src"] = name
            r = route(rec)
            if name == "xxqg":
                kw = route_kw(rec)
                if kw:
                    bucket[kw].append(rec)
                    stats[kw] += 1
                    n_ok += 1
                    continue
            if r:
                bucket[r].append(rec)
                stats[r] += 1
                n_ok += 1
        print("源 %-12s 解析 %6d 条，可用 %6d 条" % (name, n_in, n_ok))

    # 合并
    report = collections.Counter()
    total_added = 0
    for key, recs in bucket.items():
        ind, cat, cat_name = key
        if ind not in industries:
            print("!! 行业不存在，跳过:", ind)
            continue
        d = industries[ind]
        cats = d.setdefault("categories", [])
        target = next((c for c in cats if c["id"] == cat), None)
        if target is None:
            target = {"id": cat, "name": cat_name, "questions": []}
            cats.append(target)
        existing_norms = set()
        for q in target["questions"]:
            existing_norms.add(dedupe_key(q.get("q", "")))
        # 全行业查重：其他分类已有的也跳过（保留旧题）
        all_norms = set()
        for c in cats:
            if c is not target:
                for q in c["questions"]:
                    all_norms.add(dedupe_key(q.get("q", "")))
        seen = set()
        added = []
        for rec in recs:
            key = dedupe_key(rec["q"])
            if not key or key in seen or key in existing_norms or key in all_norms:
                continue
            seen.add(key)
            mq = to_site_question(rec)
            if mq:
                added.append(mq)
        target["questions"].extend(added)
        total_added += len(added)
        report[(ind, target["name"])] = len(added)

    # 写回 + 校验
    for d in industries.values():
        path = os.path.join(ROOT, "data", d["id"] + ".json")
        json.dump(d, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print("\n===== 合并结果 =====")
    per_ind = collections.Counter()
    for (ind, name), n in sorted(report.items()):
        print("  %-8s %-28s +%d" % (ind, name, n))
        per_ind[ind] += n
    print("\n行业新增合计：")
    grand = 0
    for ind, n in per_ind.most_common():
        d = industries[ind]
        tot = sum(len(c["questions"]) for c in d["categories"])
        grand += tot
        print("  %-10s 本轮 +%-6d 行业总数 %d" % (ind, n, tot))
    print("全站总题数:", grand)


def cat_name(cat_id, recs):
    # 分类名在 route() 已决定，这里兜底
    return cat_id


_PUNCT = set('，。、；：？！…—（）()【】[]{}""\'\'\'·,.:;?!'"'"'" 　\t\n\r')

def dedupe_key(q):
    s = "".join(ch for ch in str(q) if ch not in _PUNCT and not ch.isspace())
    return s[:50]


def to_site_question(rec):
    import random
    t = rec["type"]
    out = {"type": t, "q": rec["q"][:300], "explain": rec.get("explain", "")[:800] or "（暂无解析）"}
    if t in ("single", "multi"):
        opts = [o[:200] for o in rec["options"]]
        ans = rec["answer"] if t == "single" else rec["answer"]
        # 打乱选项
        order = list(range(len(opts)))
        random.shuffle(order)
        remap = {old: new for new, old in enumerate(order)}
        out["options"] = [opts[i] for i in order]
        if t == "single":
            out["answer"] = remap[ans]
        else:
            out["answers"] = sorted(remap[a] for a in ans)
    elif t == "judge":
        out["answer"] = bool(rec["answer"])
    else:
        return None
    return out


if __name__ == "__main__":
    main()
